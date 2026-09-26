"""
Excel export that reproduces the legacy workbook sheet for sheet:
"LISTAGEM DE SERVIÇOS PRESTADOS" and "RESUMO DE PGTO POR APLICADOR".

Column widths, fonts, fills, borders, number formats and row heights were read
off the reference files so a downloaded export drops straight into the routine
the team already has.
"""
from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from apps.payroll.schedule import fortnight_label
from apps.payroll.summary import SummaryGroup

# Theme colours of the reference workbook (Office "Azul, Ênfase 5 / Ênfase 1").
HEADER_BLUE = "8EAADB"  # block titles and column headers
TOTAL_BLUE = "9DC3E6"  # the "Total" row closing each summary block
WHITE = "FFFFFF"

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_LEFT = Border(left=THIN, top=THIN, bottom=THIN)
TOTAL_MIDDLE = Border(top=THIN, bottom=THIN)
TOTAL_RIGHT = Border(right=THIN, top=THIN, bottom=THIN)
TITLE_BORDER = Border(bottom=THIN)

CALIBRI_10 = Font(name="Calibri", size=10)
CALIBRI_11 = Font(name="Calibri", size=11)
CALIBRI_LIGHT_10 = Font(name="Calibri Light", size=10)
BOLD_9 = Font(name="Calibri", size=9, bold=True)
BOLD_10 = Font(name="Calibri", size=10, bold=True)
BOLD_11 = Font(name="Calibri", size=11, bold=True)

LEFT = Alignment(horizontal="left")
CENTER = Alignment(horizontal="center")
MIDDLE = Alignment(horizontal="center", vertical="center")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

GENERAL = "General"
DATE = "mm-dd-yy"  # builtin format 14, shown as dd/mm/aaaa by a pt-BR Excel
DECIMAL = "0.00"
LISTING_MONEY = '_-"R$"\\ * #,##0.00_-;\\-"R$"\\ * #,##0.00_-;_-"R$"\\ * "-"??_-;_-@_-'
SUMMARY_MONEY = '_-"R$"* #,##0.00_-;\\-"R$"* #,##0.00_-;_-"R$"* "-"??_-;_-@_-'


@dataclass(frozen=True)
class Column:
    """One spreadsheet column: its header and how every data cell below it looks."""

    header: str
    width: float
    font: Font
    alignment: Alignment
    number_format: str = GENERAL


LISTING_COLUMNS = [
    Column("NOME APLICADOR", 41, CALIBRI_10, LEFT),
    Column("VALOR LÍQUIDO", 10.77734375, CALIBRI_10, CENTER, LISTING_MONEY),
    Column("APLICADOR / ORIENTADOR", 26.6640625, CALIBRI_11, CENTER),
    Column("DATA DA PROVA / ATIVIDADE", 11.6640625, CALIBRI_10, MIDDLE, DATE),
    Column("PROVA/EVENTO", 42.33203125, CALIBRI_10, CENTER),
    Column("SEGMENTO", 18.88671875, CALIBRI_10, MIDDLE),
    Column("SETOR SOLICITANTE", 13.77734375, CALIBRI_LIGHT_10, MIDDLE),
    Column("UNIDADE ESCOLAR DA APLICAÇÃO", 15.77734375, CALIBRI_LIGHT_10, MIDDLE),
    Column("DATA PAGAMENTO", 11.77734375, CALIBRI_10, MIDDLE, DATE),
    Column("EMPRESA PAGADORA", 17.44140625, CALIBRI_LIGHT_10, MIDDLE),
    Column("OBSERVAÇÕES", 56.88671875, CALIBRI_LIGHT_10, CENTER),
    ]
LISTING_HEADER_HEIGHT = 41.4
LISTING_ROW_HEIGHT = 14.4

SUMMARY_COLUMNS = [
    Column("DATA DE PAGAMENTO", 10.6640625, CALIBRI_11, MIDDLE, DATE),
    Column("APLICADOR", 34, CALIBRI_10, LEFT),
    Column("EMPRESA", 12.44140625, CALIBRI_10, CENTER),
    Column("VALOR BRUTO (RPA)", 12.44140625, CALIBRI_11, MIDDLE, DECIMAL),
    Column("INSS", 12.21875, CALIBRI_11, MIDDLE, SUMMARY_MONEY),
    Column("ISS", 11.44140625, CALIBRI_11, MIDDLE, SUMMARY_MONEY),
    Column("IR", 7.109375, CALIBRI_11, MIDDLE, SUMMARY_MONEY),
    Column("VALOR LÍQUIDO A PAGAR", 15, CALIBRI_11, MIDDLE, SUMMARY_MONEY),
    Column("AJUSTE R$ 0,01 - ARREDONDAMENTO RPA", 17.88671875, CALIBRI_11, MIDDLE),
    Column("VALOR FINAL A PAGAR", 17.21875, CALIBRI_11, MIDDLE, SUMMARY_MONEY),
    Column("CONSISTÊNCIA", 14.44140625, CALIBRI_11, MIDDLE),
]
SUMMARY_HEADER_HEIGHT = 36
# 1-indexed summary columns the "Total" band spans, once past the merged A:C label.
TOTALS_FIRST_COLUMN = 4
ADJUSTMENT_COLUMN = 9
CONSISTENCY_COLUMN = 11


def _apply_widths(sheet: Worksheet, columns: list[Column]) -> None:
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column.width

def _write_header(sheet: Worksheet, row: int, columns: list[Column], font: Font, height: float) -> None:
    sheet.row_dimensions[row].height = height
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=row, column=index, value=column.header)
        cell.font = font
        cell.alignment = HEADER_ALIGN
        cell.fill = PatternFill("solid", fgColor=HEADER_BLUE)
        cell.border = BOX


def _write_row(sheet: Worksheet, row: int, columns: list[Column], values: list, height: float | None = None) -> None:
    if height:
        sheet.row_dimensions[row].height = height
    for index, (column, value) in enumerate(zip(columns, values), start=1):
        cell = sheet.cell(row=row, column=index, value=value)
        cell.font = column.font
        cell.alignment = column.alignment
        cell.number_format = column.number_format
        cell.fill = PatternFill("solid", fgColor=WHITE)
        cell.border = BOX

def _build_listing(sheet: Worksheet, entries) -> None:
    sheet.sheet_view.zoomScale = 70
    _apply_widths(sheet, LISTING_COLUMNS)
    _write_header(sheet, 1, LISTING_COLUMNS, BOLD_10, LISTING_HEADER_HEIGHT)
    last_row = 1
    for last_row, entry in enumerate(entries, start=2):
        _write_row(
            sheet,
            last_row,
            LISTING_COLUMNS,
            [
                entry.applicator.full_name,
                entry.net_amount,
                entry.get_role_display().upper(),
                entry.activity_date,
                entry.event_name,
                entry.segment,
                entry.sector.name,
                entry.unit.name.upper(),
                entry.payment_date,
                entry.paying_company.name,
                entry.notes,
            ],
            height=LISTING_ROW_HEIGHT,
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(LISTING_COLUMNS))}{last_row}"


def _write_block_title(sheet: Worksheet, row: int, text: str) -> None:
    """The coloured band naming the fortnight, spanning the whole table."""
    last_column = len(SUMMARY_COLUMNS)
    sheet.cell(row=row, column=1, value=text)
    # Merge before styling: merging resets the formatting of the cells it swallows.
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
    for index in range(1, last_column + 1):
        cell = sheet.cell(row=row, column=index)
        cell.font = BOLD_11
        cell.alignment = CENTER
        cell.fill = PatternFill("solid", fgColor=HEADER_BLUE)
        cell.border = TITLE_BORDER


def _write_total(sheet: Worksheet, row: int, group: SummaryGroup) -> None:
    """The "Total" band: label merged over A:C, then the summed money columns."""
    sheet.cell(row=row, column=1, value="Total")
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    for index, border in enumerate((TOTAL_LEFT, TOTAL_MIDDLE, TOTAL_RIGHT), start=1):
        cell = sheet.cell(row=row, column=index)
        cell.font = CALIBRI_11
        cell.alignment = CENTER
        cell.fill = PatternFill("solid", fgColor=TOTAL_BLUE)
        cell.border = border

    net_total = group.total("net")
    consistent = abs(group.net_payable_total - net_total) < Decimal("1")
    totals = [
        group.gross_total, group.inss_total, group.iss_total, group.ir_total,
        net_total, None, net_total, "OK" if consistent else "ver",
    ]
    for index, value in enumerate(totals, start=TOTALS_FIRST_COLUMN):
        cell = sheet.cell(row=row, column=index, value=value)
        # "Consistência" closes the band in plain text; the money columns are bold.
        cell.font = CALIBRI_11 if index == CONSISTENCY_COLUMN else BOLD_11
        cell.alignment = MIDDLE
        cell.number_format = GENERAL if index in (ADJUSTMENT_COLUMN, CONSISTENCY_COLUMN) else SUMMARY_MONEY
        cell.fill = PatternFill("solid", fgColor=TOTAL_BLUE)
        cell.border = BOX


def _build_summary(sheet: Worksheet, groups: list[SummaryGroup]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 110
    _apply_widths(sheet, SUMMARY_COLUMNS)
    # Two fortnights of different years share a name, so date them when that happens.
    labels = [fortnight_label(group.payment_date) for group in groups]
    with_year = len(set(labels)) != len(labels)
    row = 1
    for group in groups:
        _write_block_title(sheet, row, fortnight_label(group.payment_date, with_year))
        row += 1
        _write_header(sheet, row, SUMMARY_COLUMNS, BOLD_9, SUMMARY_HEADER_HEIGHT)
        row += 1
        for entry_row in group.rows:
            _write_row(
                sheet,
                row,
                SUMMARY_COLUMNS,
                [
                    entry_row.payment_date, entry_row.applicator_name, entry_row.company_name,
                    entry_row.gross, entry_row.inss, entry_row.iss, entry_row.ir,
                    entry_row.net, None, entry_row.net,
                    "OK" if entry_row.is_consistent else "ver",
                ],
            )
            row += 1
        _write_total(sheet, row, group)
        row += 2  # the blank line that separates two blocks

def build_workbook(entries, groups: list[SummaryGroup]) -> bytes:
    workbook = Workbook()
    listing = workbook.active
    listing.title = "LISTAGEM DE SERVIÇOS PRESTADOS"
    _build_listing(listing, entries)
    _build_summary(workbook.create_sheet("RESUMO DE PGTO POR APLICADOR"), groups)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
